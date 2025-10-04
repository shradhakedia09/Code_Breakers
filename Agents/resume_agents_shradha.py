import streamlit as st
import pandas as pd
import json
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
from openai import OpenAI
import pdfplumber
from pdf2image import convert_from_bytes
import pytesseract
from PIL import Image
import docx
import time

# -------------------------------------
# 🔑 Setup
# -------------------------------------
client = OpenAI(api_key=st.secrets["OPENAI_API_KEY"])

# -------------------------------------
# 🧩 Helper Functions
# -------------------------------------
def extract_text_from_pdf(uploaded_file):
    """Extract text from PDF — supports both text-based and scanned PDFs."""
    text = ""

    # 1️⃣ Try text-based extraction (pdfplumber)
    try:
        with pdfplumber.open(uploaded_file) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except Exception:
        pass

    # 2️⃣ Fallback: OCR for scanned/image PDFs
    if not text.strip():
        try:
            uploaded_file.seek(0)
            images = convert_from_bytes(uploaded_file.read())
            for img in images:
                text += pytesseract.image_to_string(img)
        except Exception as e:
            text = f"Error during OCR: {e}"

    return text.strip()


def extract_text_from_docx(uploaded_file):
    """Extract text from Word documents."""
    doc = docx.Document(uploaded_file)
    return "\n".join([para.text for para in doc.paragraphs])


def analyze_resume_with_llm(role, resume_text):
    """Send resume text to LLM for structured analysis and scoring."""
    prompt = f"""
You are an expert HR recruiter assistant.  
A candidate is applying for the role of **{role}**.  
Below is their resume text.  

1. Extract key information:
   - Age (approx if not given)
   - Education / College
   - Key Skills
   - Projects / Experience
   - Certifications (if any)

2. Score each category (0–10) based on relevance to the given role.
3. For each category, explain how relevant the candidate’s background is to the role of "{role}".
   Assign a relevance score (0–10) for each category based on how closely it matches the job requirements.

4. Compute the overall weighted score using this weightage:
   - Skills: 40%
   - Projects: 30%
   - Education: 20%
   - Certifications: 10%

5. Return only valid JSON in this format:
   {{
      "weighted_average": <float>,
      "verdict": "PASS" or "FAIL",
      "reasoning": "<short explanation>"
   }}
    """

    try:
        response = client.responses.create(
            model="gpt-4o-mini",
            input=[
                {"role": "system", "content": "You are an HR assistant. Return valid JSON only."},
                {"role": "user", "content": prompt + "\nResume:\n" + resume_text[:8000]},
            ],
            response_format={"type": "json_object"},
        )
        return response.output_text
    except Exception as e:
        return json.dumps({
            "weighted_average": 0,
            "verdict": "FAIL",
            "reasoning": f"Error during analysis: {str(e)}"
        })


# -------------------------------------
# 🚀 Streamlit App
# -------------------------------------
st.set_page_config(page_title="AI Resume Screener", page_icon="🤖", layout="centered")
st.title("🤖 AI Resume Screening Agent")

user_type = st.radio("Who are you?", ["Applicant", "HR Manager"])

# -------------------------------------
# 👤 Applicant Mode
# -------------------------------------
if user_type == "Applicant":
    role = st.text_input("Enter the role you want to apply for:")
    vacancies = {
        "Data Analyst": True,
        "Software Engineer": True,
        "Product Manager": True,
        "UI/UX Designer": True,
        "AI Research Engineer": True
    }

    if role:
        if vacancies.get(role, True):
            uploaded_file = st.file_uploader("Upload your resume (PDF or DOCX)", type=["pdf", "docx"])
            if uploaded_file:
                if uploaded_file.name.endswith(".pdf"):
                    text = extract_text_from_pdf(uploaded_file)
                else:
                    text = extract_text_from_docx(uploaded_file)

                st.write(f"📄 Extracted text length: {len(text)}")
                st.text_area("Preview of extracted text:", text[:1000])

                if text.strip():
                    with st.spinner("Analyzing your resume..."):
                        llm_output = analyze_resume_with_llm(role, text)
                        try:
                            parsed_output = json.loads(llm_output)
                            st.subheader("📊 Resume Screening Result")
                            st.json(parsed_output)

                            score = parsed_output.get("weighted_average", 0)
                            if score >= 6:
                                st.success("✅ You passed the resume screening!")
                            else:
                                st.error("❌ Sorry, you did not pass the screening.")
                        except json.JSONDecodeError:
                            st.error("⚠️ Could not parse the response from the AI.")
                else:
                    st.warning("⚠️ Could not extract text from your file.")
        else:
            st.warning(f"No current vacancies for '{role}'.")

# -------------------------------------
# 🧑‍💼 HR Manager Mode
# -------------------------------------
elif user_type == "HR Manager":
    st.subheader("📂 Bulk Resume Screening for HR Managers")
    role = st.text_input("🧾 Enter the role you are hiring for:")
    uploaded_files = st.file_uploader(
        "📎 Upload multiple resumes (PDF or DOCX)",
        type=["pdf", "docx"],
        accept_multiple_files=True
    )

    if uploaded_files and role:
        total_files = len(uploaded_files)
        results = []
        failed_files = []

        st.info(f"Processing {total_files} resumes for the role of '{role}'...")
        progress = st.progress(0)
        status_text = st.empty()
        start_time = time.time()

        for i, file in enumerate(uploaded_files):
            with st.spinner(f"Analyzing {file.name}..."):
                if file.name.endswith(".pdf"):
                    text = extract_text_from_pdf(file)
                else:
                    text = extract_text_from_docx(file)

                if text.strip():
                    llm_output = analyze_resume_with_llm(role, text)
                    try:
                        parsed = json.loads(llm_output)
                        parsed["filename"] = file.name
                        results.append(parsed)
                    except json.JSONDecodeError:
                        st.warning(f"⚠️ Could not parse {file.name}. Added to retry list.")
                        failed_files.append(file)
                else:
                    st.warning(f"⚠️ Empty or unreadable resume: {file.name}")
                    failed_files.append(file)

            progress.progress((i + 1) / total_files)
            status_text.text(f"✅ Processed {i + 1}/{total_files} resumes.")

        # 🔁 Retry Failed Resumes
        if failed_files:
            st.warning(f"⚠️ {len(failed_files)} resumes failed to parse.")
            if st.button("🔁 Retry Failed Resumes"):
                st.info("Retrying failed resumes...")
                retry_results = []
                final_failed_files = []

                for file in failed_files:
                    with st.spinner(f"Re-analyzing {file.name}..."):
                        if file.name.endswith(".pdf"):
                            text = extract_text_from_pdf(file)
                        else:
                            text = extract_text_from_docx(file)

                        if text.strip():
                            llm_output = analyze_resume_with_llm(role, text)
                            try:
                                parsed = json.loads(llm_output)
                                parsed["filename"] = file.name
                                retry_results.append(parsed)
                                st.success(f"✅ {file.name} parsed successfully on retry.")
                            except json.JSONDecodeError:
                                st.error(f"❌ Still could not parse {file.name}. Added as failed entry.")
                                retry_results.append({
                                    "filename": file.name,
                                    "weighted_average": 0,
                                    "verdict": "FAIL",
                                    "reasoning": "Resume could not be parsed after multiple attempts."
                                })
                                final_failed_files.append(file)
                        else:
                            st.error(f"❌ {file.name} unreadable (no text). Added as failed entry.")
                            retry_results.append({
                                "filename": file.name,
                                "weighted_average": 0,
                                "verdict": "FAIL",
                                "reasoning": "Resume text was empty or unreadable."
                            })
                            final_failed_files.append(file)

                results.extend(retry_results)
                failed_files = final_failed_files
                st.success(f"✅ Added {len(retry_results)} retried resumes to the results!")

        # 🧾 Final Results
        if results:
            df = pd.DataFrame(results)
            df = df.sort_values(by="weighted_average", ascending=False)

            st.success("✅ Screening complete!")
            st.dataframe(df[["filename", "weighted_average", "verdict", "reasoning"]])

            # 🟩 Excel export with color coding + Top 5 Highlight
            buffer = BytesIO()
            df.to_excel(buffer, index=False, sheet_name="Results")
            buffer.seek(0)
            wb = load_workbook(buffer)
            ws = wb.active

            green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
            gold_fill = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
            bold_font = Font(bold=True)

            headers = [cell.value for cell in ws[1]]
            verdict_col = headers.index("verdict") + 1 if "verdict" in headers else None
            reasoning_col = headers.index("reasoning") + 1 if "reasoning" in headers else None

            top5_filenames = []
            if "weighted_average" in df.columns and "filename" in df.columns:
                top5_filenames = (
                    df.sort_values(by="weighted_average", ascending=False)
                    .head(5)["filename"]
                    .tolist()
                )

            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                filename = row[0].value
                verdict = row[verdict_col - 1].value if verdict_col else ""
                reasoning = row[reasoning_col - 1].value if reasoning_col else ""

                if verdict and "PASS" in str(verdict).upper():
                    fill = green_fill
                elif "unreadable" in str(reasoning).lower() or "could not be parsed" in str(reasoning).lower():
                    fill = yellow_fill
                else:
                    fill = red_fill

                for cell in row:
                    cell.fill = fill

                if filename in top5_filenames:
                    for cell in row:
                        cell.fill = gold_fill
                        cell.font = bold_font

            final_buffer = BytesIO()
            wb.save(final_buffer)
            final_buffer.seek(0)

            st.download_button(
                label="🏆 Download Color-Coded Excel (Top 5 Highlighted)",
                data=final_buffer,
                file_name=f"resume_screening_results_{role.replace(' ', '_')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

        if not results and not failed_files:
            st.error("❌ No valid resumes were processed.")
